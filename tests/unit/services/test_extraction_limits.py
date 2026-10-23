"""Memory and size limits of document extraction (security audit 2026-10-24, M3).

python-docx and python-pptx hold every XML part as a whole tree (about 15 MB of memory per MB of XML; openpyxl 6 to
13), media costs about twice its size: the archive guard keeps a separate, smaller budget for the parts whose
content starts like XML, whatever their name. The text one document may yield is capped as well.
"""

from __future__ import annotations

import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import pymupdf
from docx import Document

from app.core.config import Settings
from app.core.document_limits import validate_ooxml_archive
from app.core.errors import DocumentLimitError
from app.services.ingestion import extract as extraction

KIB = 1024
PNG_HEAD = b"\x89PNG\r\n\x1a\n"
JPEG_HEAD = b"\xff\xd8\xff\xe0"
XML_BODY = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>'


def write_zip(path: Path, parts: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


class XmlBudgetTests(unittest.TestCase):
    def setUp(self) -> None:
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.dir = Path(temp.name)

    def check(self, parts: dict[str, bytes], *, total: int = 64 * KIB, xml: int = 4 * KIB) -> None:
        path = self.dir / "package.docx"
        write_zip(path, parts)
        validate_ooxml_archive(path, max_uncompressed_bytes=total, max_xml_bytes=xml, max_entries=50,
                               max_compression_ratio=100)

    def test_media_counts_only_against_the_total(self) -> None:
        media = os.urandom(20 * KIB)
        self.check({"word/document.xml": XML_BODY.encode(), "word/media/image1.png": PNG_HEAD + media,
                    "word/media/image2.jpeg": JPEG_HEAD + media})
        with self.assertRaises(DocumentLimitError):
            self.check({"word/media/image1.png": PNG_HEAD + media}, total=16 * KIB)

    def test_xml_parts_count_against_the_xml_budget_whatever_their_name_or_encoding(self) -> None:
        big = XML_BODY.replace("<w:body/>", "<w:body>" + "<w:p/>" * 1_000 + "</w:body>")
        shapes = {
            "named xml": ("word/document.xml", big.encode()),
            "named like a picture": ("word/media/image1.png", big.encode()),
            "utf-8 byte-order mark": ("word/document.bin", b"\xef\xbb\xbf" + big.encode()),
            "utf-16": ("word/document.bin", big.encode("utf-16")),
            "utf-32 big endian": ("word/document.bin", b"\x00\x00\xfe\xff" + big.encode("utf-32-be")),
            "long white space first": ("word/document.bin", b" \n" * 100 + big.encode()),
        }
        for name, (part, data) in shapes.items():
            with self.subTest(name), self.assertRaises(DocumentLimitError):
                self.check({part: data})
        self.check({"word/document.xml": XML_BODY.encode()})  # within the budget

    def test_the_defaults_admit_a_picture_heavy_document(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            defaults = Settings(environment="test")
        self.assertEqual((defaults.max_ooxml_uncompressed_bytes, defaults.max_ooxml_xml_bytes,
                          defaults.max_extracted_chars), (64 * KIB * KIB, 12 * KIB * KIB, 10_000_000))
        path = self.dir / "photos.docx"
        document = Document()
        document.add_paragraph("Ảnh hiện trường")
        document.save(path)
        with zipfile.ZipFile(path, "a", zipfile.ZIP_STORED) as archive:
            # Pictures do not compress: 14 MB of them is more than the XML budget and still accepted.
            archive.writestr("word/media/photo1.jpeg", JPEG_HEAD + os.urandom(14 * KIB * KIB))
        validate_ooxml_archive(path, max_uncompressed_bytes=defaults.max_ooxml_uncompressed_bytes,
                               max_xml_bytes=defaults.max_ooxml_xml_bytes, max_entries=defaults.max_ooxml_entries,
                               max_compression_ratio=defaults.max_ooxml_compression_ratio)
        self.assertEqual(extraction.extract_sections(path, path.name)[0].text, "Ảnh hiện trường")


class ExtractedCharacterLimitTests(unittest.TestCase):
    def setUp(self) -> None:
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.dir = Path(temp.name)

    def test_plain_text_is_measured_before_cleaning(self) -> None:
        path = self.dir / "notes.txt"
        path.write_text("a  " * 400, encoding="utf-8")  # 1,200 characters, 799 after cleaning
        with patch.object(extraction.settings, "max_extracted_chars", 1_000), \
                patch.object(extraction, "clean_text") as cleaner, self.assertRaises(DocumentLimitError):
            extraction.extract_sections(path, path.name)
        cleaner.assert_not_called()
        with patch.object(extraction.settings, "max_extracted_chars", 1_200):
            self.assertEqual(len(extraction.extract_sections(path, path.name)[0].text), 799)

    def test_pdf_stops_at_the_page_that_crosses_the_limit(self) -> None:
        path = self.dir / "book.pdf"
        document = pymupdf.open()
        for number in range(5):
            document.new_page().insert_text((72, 72), f"Trang {number + 1}: " + "chữ " * 20)
        document.save(path)
        document.close()
        pages = [len(section.text) for section in extraction.extract_pdf(path)]
        seen: list[int] = []
        original = extraction._pdf_page_content

        def count(page: object) -> tuple[str, dict[str, object]]:
            seen.append(1)
            return original(page)

        with patch.object(extraction.settings, "max_extracted_chars", sum(pages[:2]) + 1), \
                patch.object(extraction, "_pdf_page_content", side_effect=count), \
                self.assertRaises(DocumentLimitError):
            extraction.extract_pdf(path)
        self.assertEqual(len(seen), 3)
        with patch.object(extraction.settings, "max_extracted_chars", sum(pages)):
            self.assertEqual([len(section.text) for section in extraction.extract_pdf(path)], pages)

    def test_every_format_is_measured_after_extraction(self) -> None:
        path = self.dir / "guide.docx"
        document = Document()
        document.add_paragraph("Quy trình kiểm soát rủi ro gồm bốn bước.")
        document.save(path)
        length = len(extraction.extract_sections(path, path.name)[0].text)
        with patch.object(extraction.settings, "max_extracted_chars", length - 1), \
                self.assertRaises(DocumentLimitError):
            extraction.extract_sections(path, path.name)
        with patch.object(extraction.settings, "max_extracted_chars", length):
            self.assertEqual(len(extraction.extract_sections(path, path.name)[0].text), length)


if __name__ == "__main__":
    unittest.main()
