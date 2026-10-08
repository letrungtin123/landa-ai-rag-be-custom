from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from app.services.ingestion.chunking import build_chunks, split_text
from app.services.ingestion.extract import (
    ExtractedSection,
    STRUCTURED_EXTRACTION_VERSION,
    _render_structured_table,
    extract_docx,
    extract_pdf,
    extract_pptx,
    extract_xlsx,
)
from app.source_structure import analyze_source_structure


class StructuredSourceExtractionTests(unittest.TestCase):
    def test_oversized_table_splits_only_between_complete_rows(self) -> None:
        table = "[TABLE]\n" + "\n".join(
            f"Row {index}: Bậc {index:02d} | Năng lực {index} | Mô tả vận hành {index}"
            for index in range(1, 9)
        )

        chunks = split_text(table, max_chars=150, overlap_chars=20)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.startswith("[TABLE]\nRow ") for chunk in chunks))
        rows = [line for chunk in chunks for line in chunk.splitlines() if line.startswith("Row ")]
        self.assertEqual(len(rows), 8)
        self.assertTrue(all(" | " in row for row in rows))

    def test_chunk_metadata_marks_only_the_chunk_that_contains_a_table(self) -> None:
        section = ExtractedSection(
            text=("Giới thiệu ngắn.\n\n[TABLE]\nRow 1: Bậc | Năng lực\n"
                  "Row 2: Bậc 04 | Vietnam Know-how\n\nKết luận ngắn."),
            page=1,
            metadata={
                "extraction_version": STRUCTURED_EXTRACTION_VERSION,
                "content_kinds": ["text", "table"],
                "table_count": 1,
            },
        )

        with patch("app.main.settings.chunk_max_chars", 80), \
                patch("app.main.settings.chunk_overlap_chars", 0):
            chunks = build_chunks([section])
        table_chunks = [chunk for chunk in chunks if "[TABLE]" in chunk["content"]]
        non_table_chunks = [chunk for chunk in chunks if "[TABLE]" not in chunk["content"]]

        self.assertEqual(len(table_chunks), 1)
        self.assertEqual(table_chunks[0]["metadata"]["table_count"], 1)
        self.assertIn("table", table_chunks[0]["metadata"]["content_kinds"])
        self.assertTrue(all(chunk["metadata"]["table_count"] == 0 for chunk in non_table_chunks))
        self.assertTrue(all("table" not in chunk["metadata"]["content_kinds"] for chunk in non_table_chunks))
        revisions = {chunk["metadata"]["source_evidence_revision"] for chunk in chunks}
        self.assertEqual(len(revisions), 1)
        self.assertTrue(all(
            chunk["metadata"]["structured_evidence_contract_version"] == "source-evidence-propagation-v1"
            for chunk in chunks
        ))

    def test_table_renderer_preserves_sparse_spreadsheet_coordinates(self) -> None:
        rendered = _render_structured_table(
            [["Control", None, "Owner"], ["Isolation", "", "Supervisor"]],
            coordinate_labels=True,
        )

        self.assertIn("A1=Control", rendered)
        self.assertIn("C1=Owner", rendered)
        self.assertIn("A2=Isolation", rendered)
        self.assertIn("C2=Supervisor", rendered)
        self.assertNotIn("B1=", rendered)

    def test_native_table_renderer_preserves_blank_cell_position(self) -> None:
        rendered = _render_structured_table([
            ["Hazard", "Control", "Owner"],
            ["Chemical", "", "HSE lead"],
        ])

        self.assertIn("Row 2: Chemical |  | HSE lead", rendered)

    def test_docx_preserves_heading_paragraph_and_table_document_order(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "structured.docx"
            document = Document()
            document.add_heading("Hazard identification", level=1)
            document.add_paragraph("Review the work area before starting.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Hazard"
            table.cell(0, 1).text = "Control"
            table.cell(1, 0).text = "Energy"
            table.cell(1, 1).text = "Isolation"
            document.add_paragraph("Escalate any uncontrolled risk.")
            document.save(path)

            section = extract_docx(path)[0]

        self.assertLess(section.text.index("Review the work area"), section.text.index("[TABLE]"))
        self.assertLess(section.text.index("[TABLE]"), section.text.index("Escalate any uncontrolled risk"))
        self.assertEqual(section.metadata["table_count"], 1)
        self.assertEqual(section.metadata["heading_candidates"], [
            {"title": "Hazard identification", "level": 1},
        ])

    def test_pdf_uses_layout_font_signal_for_heading_metadata(self) -> None:
        import pymupdf

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "structured.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((72, 72), "Emergency response", fontsize=20)
            page.insert_text((72, 120), "Follow the approved alarm and evacuation procedure.", fontsize=10)
            document.save(path)
            document.close()

            section = extract_pdf(path)[0]

        self.assertIn("Emergency response", section.text)
        self.assertEqual(section.metadata["extraction_version"], STRUCTURED_EXTRACTION_VERSION)
        self.assertIn("Emergency response", {
            item["title"] for item in section.metadata["heading_candidates"]
        })

    def test_pptx_preserves_title_and_table_as_distinct_structure(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "structured.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "PPE inspection"
            shape = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5))
            shape.table.cell(0, 0).text = "Item"
            shape.table.cell(0, 1).text = "Check"
            shape.table.cell(1, 0).text = "Helmet"
            shape.table.cell(1, 1).text = "Damage"
            presentation.save(path)

            section = extract_pptx(path)[0]

        self.assertIn("PPE inspection", section.text)
        self.assertIn("[TABLE]", section.text)
        self.assertEqual(section.metadata["table_count"], 1)
        self.assertEqual(section.metadata["heading_candidates"], [
            {"title": "PPE inspection", "level": 1},
        ])

    def test_xlsx_preserves_cell_coordinates_and_formula_text(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "structured.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Risk matrix"
            sheet["A1"] = "Likelihood"
            sheet["C1"] = "Risk"
            sheet["A2"] = 2
            sheet["B2"] = 3
            sheet["C2"] = "=A2*B2"
            workbook.save(path)

            section = extract_xlsx(path)[0]

        self.assertIn("A1=Likelihood", section.text)
        self.assertIn("C1=Risk", section.text)
        self.assertIn("C2==A2*B2", section.text)
        self.assertEqual(section.section, "Risk matrix")

    def test_structured_headings_are_used_before_heuristic_uppercase_detection(self) -> None:
        structure = analyze_source_structure([
            ExtractedSection(
                text="A sentence-case chapter title\nDetailed body text ends with a period.",
                page=1,
                metadata={
                    "heading_candidates": [{"title": "A sentence-case chapter title", "level": 1}],
                },
            ),
        ])

        self.assertEqual(structure["structure_source"], "heading_inferred")
        self.assertEqual(structure["nodes"][0]["title"], "A sentence-case chapter title")
        self.assertIn("STRUCTURED_HEADINGS_USED", structure["warnings"])

    def test_chunk_metadata_keeps_extraction_and_structure_provenance(self) -> None:
        section = ExtractedSection(
            text="1. Hazard controls\nApply the approved isolation procedure.",
            page=1,
            metadata={
                "extraction_version": STRUCTURED_EXTRACTION_VERSION,
                "content_kinds": ["text"],
            },
        )

        chunks = build_chunks([section])

        self.assertEqual(chunks[0]["metadata"]["extraction_version"], STRUCTURED_EXTRACTION_VERSION)
        self.assertEqual(chunks[0]["metadata"]["content_kinds"], ["text"])
        self.assertIn("parser_version", chunks[0]["metadata"])

    def test_structured_evidence_change_materializes_a_new_source_revision(self) -> None:
        base_metadata = {
            "extraction_version": STRUCTURED_EXTRACTION_VERSION,
            "content_kinds": ["text", "image"],
            "visual_regions": [{
                "region_kind": "embedded_image",
                "asset_revision": "a" * 64,
                "locator": {"page": 1, "bbox_normalized": [0.1, 0.2, 0.8, 0.9]},
                "observation": {"status": "unreviewed", "facts": []},
                "inference": {"status": "not_performed", "claims": []},
            }],
        }
        original = build_chunks([ExtractedSection(
            text="Review the control hierarchy diagram.", page=1, metadata=base_metadata,
        )])
        unchanged = build_chunks([ExtractedSection(
            text="Review the control hierarchy diagram.", page=1, metadata=base_metadata,
        )])
        enriched_metadata = {
            **base_metadata,
            "visual_regions": [{
                **base_metadata["visual_regions"][0],
                "observation": {
                    "status": "observed",
                    "facts": ["The hierarchy contains five ordered control levels."],
                },
            }],
        }
        enriched = build_chunks([ExtractedSection(
            text="Review the control hierarchy diagram.", page=1, metadata=enriched_metadata,
        )])

        original_revision = original[0]["metadata"]["source_evidence_revision"]
        self.assertEqual(
            original_revision,
            unchanged[0]["metadata"]["source_evidence_revision"],
        )
        self.assertNotEqual(
            original_revision,
            enriched[0]["metadata"]["source_evidence_revision"],
        )


if __name__ == "__main__":
    unittest.main()
